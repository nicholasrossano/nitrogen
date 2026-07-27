from io import BytesIO

import openpyxl
import pytest

from app.domain.energy.assessments.risk_assessment import RiskAssessment
from app.assessments.utils import make_build_item


def test_risk_assessment_stage_defs_follow_expected_flow():
    assessment = RiskAssessment()

    assert assessment.definition.id == "risk_assessment"
    assert assessment.definition.export_format == "xlsx"
    assert [stage.id for stage in assessment.stage_defs] == [
        "categories",
        "risks",
        "mitigations",
        "results",
    ]
    assert assessment.stage_defs[-1].title == "Results"
    assert assessment.stage_defs[-1].widget == "risk_register_results"
    assert assessment.manifest.adapter_bindings == {"research_source": "retrieval"}
    assert assessment.manifest.produced_outputs == ["risk_register"]


@pytest.mark.asyncio
async def test_compute_stage_builds_register_and_normalizes_ratings(monkeypatch):
    assessment = RiskAssessment()
    category = make_build_item({
        "label": "Institutional Capacity",
        "description": "Implementing agency and partner capability.",
        "why_it_matters": "Several delivery partners are expected.",
    })
    risk = make_build_item({
        "title": "Multiple delivery channels may create coordination delays",
        "category": "Institutional Capacity",
        "description": "Grid, mini-grid, household systems, and clean cooking channels require coordination.",
        "affected_components": "Grid; mini-grid; household systems; clean cooking",
        "basis": "The project involves several technologies and likely implementation partners.",
        "missing_information": "Implementation roles and coordination arrangements are not yet documented.",
    })

    async def fake_rate_register(context, category_items, draft_register):
        return {
            "category_ratings": [{
                "category": "Institutional Capacity",
                "rating": "substantial",
                "rationale": "Coordination burden remains material.",
                "top_risks": ["Multiple delivery channels may create coordination delays"],
                "unresolved_issues": ["Implementation roles are not yet documented."],
            }],
            "risk_register": [{
                "risk_id": "R01",
                "inherent_rating": "High",
                "residual_rating": "moderate",
                "rating_rationale": "Mitigation reduces but does not eliminate coordination risk.",
                "basis_evidence": "Project has multiple delivery channels.",
                "missing_information": "Implementation roles and coordination arrangements are not yet documented.",
            }],
        }

    monkeypatch.setattr(assessment, "_rate_register", fake_rate_register)

    result = await assessment.compute_stage(
        "results",
        {
            "categories": {"data": {"items": [category]}},
            "risks": {"data": {"items": [risk]}},
            "mitigations": {
                "data": {
                    "records": {
                        risk["id"]: {
                            "mitigation": "Define agency roles and escalation protocols before launch.",
                            "owner": "Project team",
                            "timing": "Preparation",
                            "remaining_issue": "Partner capacity needs validation.",
                            "status": "Needs validation",
                        }
                    }
                }
            },
        },
        {"project_title": "Energy Access Project"},
    )

    row = result["risk_register"][0]
    assert row["risk_id"] == "R01"
    assert row["category"] == "Institutional Capacity"
    assert row["inherent_rating"] == "High"
    assert row["residual_rating"] == "Moderate"
    assert row["mitigation"] == "Define agency roles and escalation protocols before launch."
    assert result["category_ratings"][0]["rating"] == "Substantial"
    assert "Risk ID" in result["copy"]["markdown"]
    assert "risk_id" in result["copy"]["tsv"]


def test_export_xlsx_includes_register_sheets():
    assessment = RiskAssessment()
    workbook_bytes = assessment.export_xlsx({
        "risk_register": [{
            "risk_id": "R01",
            "category": "Technical Design",
            "risk_title": "Design variables may not match site conditions",
            "description": "Engineering inputs need validation.",
            "affected_components": "System design",
            "inherent_rating": "Substantial",
            "mitigation": "Complete site validation before procurement.",
            "residual_rating": "Moderate",
            "owner_status": "Project team / Needs validation",
            "basis_evidence": "Project materials reference early-stage design.",
            "missing_information": "Final site survey.",
            "rating_rationale": "Residual risk remains until survey is complete.",
        }],
        "category_ratings": [{
            "category": "Technical Design",
            "rating": "Moderate",
            "rationale": "Site validation can reduce risk.",
        }],
        "top_risks": [{
            "risk_id": "R01",
            "risk_title": "Design variables may not match site conditions",
            "why_it_matters": "Procurement could be mis-specified.",
            "mitigation_summary": "Complete site validation.",
        }],
        "unresolved_issues": [{
            "risk_id": "R01",
            "issue": "Final site survey.",
        }],
    })

    workbook = openpyxl.load_workbook(BytesIO(workbook_bytes))
    assert workbook.sheetnames == [
        "Risk Register",
        "Category Ratings",
        "Top Risks",
        "Unresolved Issues",
    ]
    register = workbook["Risk Register"]
    assert register["A1"].value == "Risk ID"
    assert register["A2"].value == "R01"
    assert register["H2"].value == "Moderate"


@pytest.mark.asyncio
async def test_bulk_mitigation_generation_populates_each_risk(monkeypatch):
    assessment = RiskAssessment()
    risk_item = make_build_item({
        "title": "Fragmented beneficiary and site data could weaken targeting and implementation planning.",
        "category": "Data Quality and Results Verification",
        "affected_components": "Beneficiary targeting; site prioritization",
        "why_it_matters": "The project depends on credible data to choose sites and beneficiaries.",
        "evidence_basis": "Project context references geospatial analysis and site planning.",
        "missing_information": "Data owners and QA process are not documented.",
    })

    async def fake_llm_json(*args, **kwargs):
        return {
            "mitigations": [{
                "source_item_id": risk_item["id"],
                "mitigation": "Define a data dictionary and QA workflow before site prioritization is finalized.",
                "owner": "Project data lead",
                "timing": "Preparation",
                "remaining_issue": "Data-sharing authority is still unconfirmed.",
                "status": "Needs validation",
            }]
        }

    monkeypatch.setattr("app.domain.energy.assessments.risk_assessment.llm_json", fake_llm_json)

    records = await assessment.enrich_records_for_stage(
        "mitigations",
        [risk_item],
        {},
        {
            "project_title": "Malawi Energy Access Project",
            "project_type": "energy_access",
            "geography": "Malawi",
            "project_description": "Uses geospatial analysis for energy access site prioritization.",
        },
    )

    assert risk_item["id"] in records
    assert records[risk_item["id"]]["mitigation"].startswith("Define a data dictionary")
    assert records[risk_item["id"]]["owner"] == "Project data lead"


def _risk_payload(title: str, category: str, **overrides):
    payload = {
        "title": title,
        "category": category,
        "affected_components": "Procurement; schedule",
        "why_it_matters": "Delay would push the critical path.",
        "evidence_basis": "Project materials reference an unconfirmed timetable.",
        "missing_information": "Approved timetable.",
        "evidence_status": "Needs evidence",
    }
    payload.update(overrides)
    return payload


CATEGORIES = [
    make_build_item({
        "label": "Sector Policy and Regulatory",
        "description": "Rules and permits",
        "why_it_matters": "Project depends on policy approvals.",
        "status": "Include",
    }),
    make_build_item({
        "label": "Technical Design and Delivery",
        "description": "Engineering and integration",
        "why_it_matters": "Delivery complexity is high.",
        "status": "Include",
    }),
]


@pytest.mark.parametrize(
    "context",
    [
        {"project_type": "energy_access", "geography": "Kenya", "target_population": "rural households"},
        {"project_type": "water_sanitation", "geography": "Peru", "target_population": "peri-urban residents"},
        {"project_type": "urban_transport", "geography": "Vietnam", "target_population": "commuters"},
        {"project_type": "land_restoration", "geography": "Brazil", "target_population": "smallholder farmers"},
    ],
)
@pytest.mark.asyncio
async def test_generate_risks_keeps_substantive_rows_across_project_demographics(monkeypatch, context):
    """Screening must not depend on rows echoing their own sector/geography terms."""
    assessment = RiskAssessment()
    proposed = [
        _risk_payload(
            "Delays in tender evaluation could postpone contract award beyond the funding window",
            "Sector Policy and Regulatory",
        ),
        _risk_payload(
            "Tariff approval that lags financial close would undermine cost recovery assumptions",
            "Sector Policy and Regulatory",
        ),
        _risk_payload(
            "Unvalidated site survey data may force redesign after procurement is committed",
            "Technical Design and Delivery",
        ),
        _risk_payload(
            "Interface gaps between parallel work packages could strand commissioning milestones",
            "Technical Design and Delivery",
        ),
    ]

    async def fake_llm_json(*args, **kwargs):
        return {"risks": proposed}

    monkeypatch.setattr("app.domain.energy.assessments.risk_assessment.llm_json", fake_llm_json)

    risks = await assessment._generate_risks(context, CATEGORIES)

    assert len(risks) == 4
    assert {risk["category"] for risk in risks} == {
        "Sector Policy and Regulatory",
        "Technical Design and Delivery",
    }
    assert all(risk["evidence_status"] == "Needs evidence" for risk in risks)


@pytest.mark.asyncio
async def test_generate_risks_drops_placeholders_and_label_only_rows(monkeypatch):
    assessment = RiskAssessment()

    async def fake_llm_json(*args, **kwargs):
        return {
            "risks": [
                _risk_payload("Needs validation", "Sector Policy and Regulatory"),
                _risk_payload("Procurement risk", "Sector Policy and Regulatory"),
                _risk_payload("Design variables to be confirmed", "Technical Design and Delivery"),
                _risk_payload(
                    "Missing hydrological baseline could invalidate sizing assumptions at final design",
                    "Technical Design and Delivery",
                    why_it_matters="",
                ),
                _risk_payload(
                    "Permit sequencing that slips past mobilization could idle the contractor",
                    "Sector Policy and Regulatory",
                ),
            ]
        }

    monkeypatch.setattr("app.domain.energy.assessments.risk_assessment.llm_json", fake_llm_json)

    risks = await assessment._generate_risks({"project_type": "energy_access", "geography": "Kenya"}, CATEGORIES)

    titles = [risk["title"] for risk in risks]
    assert titles == ["Permit sequencing that slips past mobilization could idle the contractor"]


@pytest.mark.asyncio
async def test_generate_risks_dedupes_and_caps_per_category(monkeypatch):
    assessment = RiskAssessment()

    async def fake_llm_json(*args, **kwargs):
        return {
            "risks": [
                _risk_payload(f"Escalating input costs could erode budget headroom in phase {n}", "Technical Design and Delivery")
                for n in range(1, 8)
            ] + [
                # Same statement, different casing/spacing — must collapse to one row.
                _risk_payload("Escalating input costs could erode budget headroom in phase 1", "Technical Design and Delivery"),
                _risk_payload("ESCALATING  input costs could erode budget headroom in phase 1", "Technical Design and Delivery"),
            ]
        }

    monkeypatch.setattr("app.domain.energy.assessments.risk_assessment.llm_json", fake_llm_json)

    risks = await assessment._generate_risks({"project_type": "energy_access"}, CATEGORIES)

    technical = [risk for risk in risks if risk["category"] == "Technical Design and Delivery"]
    assert len(technical) == 4
    assert len({risk["title"].lower() for risk in technical}) == 4


@pytest.mark.asyncio
async def test_thin_categories_are_topped_up_by_a_second_model_pass(monkeypatch):
    """Depth comes from another model call, never from static template text."""
    assessment = RiskAssessment()
    calls = []

    async def fake_llm_json(*args, **kwargs):
        calls.append(kwargs.get("user_msg", ""))
        if len(calls) == 1:
            return {
                "risks": [
                    _risk_payload(
                        "Tariff approval that lags financial close would undermine cost recovery",
                        "Sector Policy and Regulatory",
                    ),
                    _risk_payload(
                        "Permit sequencing that slips past mobilization could idle the contractor",
                        "Sector Policy and Regulatory",
                    ),
                ]
            }
        return {
            "risks": [
                _risk_payload(
                    "Unvalidated load profiles could size the system incorrectly at detailed design",
                    "Technical Design and Delivery",
                ),
                _risk_payload(
                    "Interface gaps between work packages could strand commissioning milestones",
                    "Technical Design and Delivery",
                ),
            ]
        }

    monkeypatch.setattr("app.domain.energy.assessments.risk_assessment.llm_json", fake_llm_json)

    risks = await assessment._generate_risks({"project_type": "energy_access", "geography": "Kenya"}, CATEGORIES)

    assert len(calls) == 2
    assert "still need more risks" in calls[1]
    assert "Technical Design and Delivery" in calls[1]
    technical = [risk for risk in risks if risk["category"] == "Technical Design and Delivery"]
    assert len(technical) == 2


@pytest.mark.asyncio
async def test_top_up_failure_keeps_first_pass_risks(monkeypatch):
    assessment = RiskAssessment()
    calls = []

    async def fake_llm_json(*args, **kwargs):
        calls.append(1)
        if len(calls) == 1:
            return {
                "risks": [
                    _risk_payload(
                        "Tariff approval that lags financial close would undermine cost recovery",
                        "Sector Policy and Regulatory",
                    ),
                ]
            }
        raise RuntimeError("model unavailable")

    monkeypatch.setattr("app.domain.energy.assessments.risk_assessment.llm_json", fake_llm_json)

    risks = await assessment._generate_risks({"project_type": "energy_access"}, CATEGORIES)

    assert len(risks) == 1
    assert risks[0]["category"] == "Sector Policy and Regulatory"
